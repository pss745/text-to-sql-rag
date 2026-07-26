import os
from datetime import datetime

from openpyxl import Workbook, load_workbook

LOG_FILE = "query_log.xlsx"


def log_query(question: str, sql: str):
    if os.path.exists(LOG_FILE):
        wb = load_workbook(LOG_FILE)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["timestamp", "question", "generated_sql"])

    ws.append([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), question, sql])
    wb.save(LOG_FILE)